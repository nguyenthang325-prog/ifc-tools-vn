"""
Bóc khối lượng (QTO) từ file IFC — đọc BaseQuantities Psets.
Output: Excel 3 sheets — Chi tiết | Tổng hợp | Tóm tắt BoQ

Cách dùng:
    python ifc_qto.py --ifc Duplex_A_20110907.ifc --discipline ALL --project Duplex --bo-sung-qto
    python ifc_qto.py --ifc file.ifc --discipline ALL --stage TKKT --project "TenDuAn"

Lưu ý: IFC phải xuất với "Export base quantities" được bật trong Revit IFC Exporter.
  Revit → File → Export → IFC → Modify Setup → check "Export base quantities"
  Không xuất lại được (file của bên khác, mô hình sinh bằng code) → thêm --bo-sung-qto:
  đo từ hình học, ghi bản sao <tên>_BoSungQTO.ifc (dùng tiếp cho wbs_ifc_mapper/model_payment).
"""
import sys
import re
import argparse
from pathlib import Path
from datetime import date
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

try:
    import ifcopenshell
    import ifcopenshell.util.element as ifc_util
except ImportError:
    print("Lỗi: pip install ifcopenshell")
    sys.exit(1)

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Lỗi: pip install openpyxl")
    sys.exit(1)

from ifc_tools_vn.config import BASE, KET_QUA_DIR, OUTPUT_DIR  # đường dẫn: xem config.py

# ── Element types theo discipline ─────────────────────────────────────────────

DISC_TYPES = {
    "KT": [
        "IfcWall", "IfcWallStandardCase", "IfcSlab", "IfcRoof",
        "IfcDoor", "IfcWindow", "IfcStair", "IfcRailing", "IfcRamp",
        "IfcCovering", "IfcFurnishingElement",
    ],
    "KC": [
        "IfcBeam", "IfcColumn", "IfcSlab",
        "IfcPlate", "IfcMember", "IfcFooting",
        "IfcWall", "IfcWallStandardCase",
    ],
    "MEP": [
        "IfcPipeSegment", "IfcPipeFitting",
        "IfcDuctSegment", "IfcDuctFitting",
        "IfcFlowSegment", "IfcFlowTerminal",
        "IfcFlowController", "IfcFlowFitting",
        "IfcLightFixture", "IfcAirTerminal",
        "IfcCableCarrierSegment",
    ],
}

TYPE_GROUP_VN = {
    "IfcWall": "Tường", "IfcWallStandardCase": "Tường",
    "IfcSlab": "Sàn/Mái", "IfcRoof": "Mái",
    "IfcBeam": "Dầm", "IfcColumn": "Cột",
    "IfcDoor": "Cửa", "IfcWindow": "Cửa sổ",
    "IfcStair": "Cầu thang", "IfcRailing": "Lan can", "IfcRamp": "Dốc",
    "IfcPlate": "Tấm KC", "IfcMember": "Thanh KC", "IfcFooting": "Móng",
    "IfcPipeSegment": "Ống nước", "IfcPipeFitting": "Phụ kiện ống nước",
    "IfcDuctSegment": "Ống gió", "IfcDuctFitting": "Phụ kiện ống gió",
    "IfcFlowSegment": "MEP", "IfcFlowTerminal": "Thiết bị MEP",
    "IfcFlowController": "Thiết bị điều khiển", "IfcFlowFitting": "Phụ kiện MEP",
    "IfcLightFixture": "Đèn chiếu sáng", "IfcAirTerminal": "Miệng gió",
    "IfcCableCarrierSegment": "Máng cáp điện",
    "IfcCovering": "Ốp lát", "IfcFurnishingElement": "Nội thất",
    "IfcReinforcingBar": "Cốt thép thanh", "IfcReinforcingMesh": "Lưới thép",
    "IfcTendon": "Cáp ứng suất trước",
}

# Heuristic đơn vị từ tên property
UNIT_HINT = {
    "volume": "m³",
    "area": "m²", "sidearea": "m²", "footprintarea": "m²", "projectedarea": "m²",
    "length": "m", "width": "m", "height": "m", "depth": "m",
    "perimeter": "m", "thickness": "m",
    "weight": "kg", "mass": "kg",
    "count": "cái", "number": "cái",
    # Cốt thép (giá trị Revit Dimensions pset — đơn vị mm)
    "diameter": "mm", "spacing": "mm", "radius": "mm",
    "quantity": "cái",
}

# Đại lượng ưu tiên khi chọn KL chính cho Tom_Tat_BoQ (số nhỏ = ưu tiên cao)
_PRIMARY_SCORE = {
    "grossvolume": 1, "netvolume": 1, "volume": 1,
    "netsidearea": 2, "grosssidearea": 2, "netarea": 2, "grossarea": 2, "area": 2,
    "length": 3, "perimeter": 3,
}


def _guess_unit(prop_name: str, length_unit: str = "m") -> str:
    p = prop_name.lower()
    for key, unit in UNIT_HINT.items():
        if key in p:
            # Đại lượng dài (length/width/height/perimeter/thickness) theo đơn vị dài
            # của project IFC. Revit IFC2x3 xuất IfcQuantityLength theo project unit (mm),
            # còn Area/Volume theo m²/m³ → chỉ thay khi hint là linear ("m").
            return length_unit if unit == "m" else unit
    return "—"


def _length_unit_symbol(ifc_file) -> str:
    """Đơn vị dài của project IFC (mm/cm/m/ft). Revit thường xuất length = mm."""
    _PREFIX = {"MILLI": "mm", "CENTI": "cm", "DECI": "dm",
               "DECA": "dam", "KILO": "km", None: "m"}
    try:
        proj = ifc_file.by_type("IfcProject")
        for u in (proj[0].UnitsInContext.Units if proj else []) or []:
            if u.is_a("IfcSIUnit") and u.UnitType == "LENGTHUNIT":
                return _PREFIX.get(u.Prefix, "m")
            if u.is_a("IfcConversionBasedUnit") and u.UnitType == "LENGTHUNIT":
                return (u.Name or "unit").lower()
    except Exception:
        pass
    return "m"


def _priority(prop_name: str) -> int:
    p = prop_name.lower()
    for key, score in _PRIMARY_SCORE.items():
        if key in p:
            return score
    return 99


# ── Excel styles ──────────────────────────────────────────────────────────────

_H_FILL  = PatternFill("solid", fgColor="1F4E79")
_S_FILL  = PatternFill("solid", fgColor="2E75B6")
_ALT     = PatternFill("solid", fgColor="DEEAF1")
_W_FILL  = PatternFill("solid", fgColor="FFF2CC")
_thin    = Side(style="thin", color="CCCCCC")
_border  = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_center  = Alignment(horizontal="center", vertical="center", wrap_text=True)
_left    = Alignment(horizontal="left",   vertical="center", wrap_text=True)
_H_FONT  = Font(bold=True, color="FFFFFF", size=10)
_B_FONT  = Font(bold=True, size=10)
_N_FONT  = Font(size=10)
_S_FONT  = Font(size=9)


def _c(ws, row, col, value=None, fill=None, font=None, align=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.border    = _border
    cell.fill      = fill or PatternFill()
    cell.font      = font or _N_FONT
    cell.alignment = align or _center
    return cell


def _col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ── IFC helpers ───────────────────────────────────────────────────────────────

def _storey(element) -> str:
    for rel in getattr(element, "ContainedInStructure", None) or []:
        container = rel.RelatingStructure
        if container.is_a("IfcBuildingStorey"):
            return container.Name or ""
    return ""


_REBAR_TYPES = {"IfcReinforcingBar", "IfcReinforcingMesh", "IfcTendon"}
_REBAR_WEIGHT_KEYS = {"weight", "mass"}
_STEEL_DENSITY = 7850  # kg/m³

# Properties hữu ích từ pset Dimensions và Structural (Revit native, đơn vị mm)
_REBAR_DIM_PROPS  = {"bar diameter", "total bar length", "bar length", "bar segment length"}
_REBAR_STR_PROPS  = {"estimated reinforcement volume", "reinforcement volume"}


def _rebar_weight_from_step(el):
    """Trọng lượng (kg) 1 thanh thép từ THUỘC TÍNH STEP của IfcReinforcingBar.

    Revit IFC2x3 KHÔNG xuất Qto_ cho cốt thép, nhưng luôn ghi CrossSectionArea (m²)
    và BarLength (mm) thẳng vào thực thể. Không có nhánh này thì mọi mô hình Revit
    có thép đều bị báo nhầm là "không có dữ liệu cốt thép".
    Đồng bộ với model_payment.py / payment_from_model.py → _rebar_weight_kg().
    """
    try:
        length = float(getattr(el, "BarLength", None) or 0)
        if length <= 0:
            return None
        length_m = length / 1000.0 if length > 100 else length   # >100 → chắc chắn là mm
        area = float(getattr(el, "CrossSectionArea", None) or 0)
        if not (1e-5 < area < 5e-3):                             # ngoài dải m² hợp lý của rebar
            area = area / 1e6 if area > 10 else 0                # mm² → m²
        if area <= 0:
            return None
        return area * length_m * _STEEL_DENSITY
    except (TypeError, ValueError):
        return None


def _extract_rebar(ifc_file, discipline: str) -> list:
    """
    Trích dữ liệu cốt thép: ưu tiên Psets (Qto_ReinforcingBarBaseQuantities + Pset_Rebar*),
    nếu Pset không cho ra trọng lượng thì tính từ thuộc tính STEP (xem _rebar_weight_from_step).
    Trả về list record, hoặc [] nếu không có entity cốt thép.
    """
    disc = discipline.upper()
    if disc not in ("KC", "ALL"):
        return []

    length_unit = _length_unit_symbol(ifc_file)
    records = []
    for ifc_type in _REBAR_TYPES:
        try:
            elements = ifc_file.by_type(ifc_type)
        except Exception:
            continue
        for el in elements:
            name   = (el.Name or "").strip() or "—"
            storey = _storey(el)
            guid   = el.GlobalId or ""
            try:
                all_psets = ifc_util.get_psets(el)
            except Exception:
                all_psets = {}

            # Thu thập Qto_*/BaseQuantities/Rebar*/Reinforc* — dữ liệu chuẩn IFC
            qtos = {k: v for k, v in all_psets.items()
                    if (k.startswith("Qto_") or k == "BaseQuantities"
                        or "Rebar" in k or "Reinforc" in k)
                    and isinstance(v, dict)}

            # Bổ sung từ Dimensions và Structural (Revit native) — chỉ lấy props hữu ích
            for pname in ("Dimensions", "Structural"):
                psdata = all_psets.get(pname)
                if not isinstance(psdata, dict):
                    continue
                allowed = _REBAR_DIM_PROPS | _REBAR_STR_PROPS
                filtered = {k: v for k, v in psdata.items() if k.lower() in allowed}
                if filtered:
                    qtos[pname] = filtered

            got_weight = False
            for pset_name, props in qtos.items():
                for prop, val in props.items():
                    if prop.lower() == "id" or val is None:
                        continue
                    try:
                        num_val = float(val)
                    except (TypeError, ValueError):
                        continue
                    pl = prop.lower()
                    if "volume" in pl:
                        # Volume cốt thép → đổi sang kg (7850 kg/m³)
                        num_val = round(num_val * _STEEL_DENSITY, 2)
                        unit = "kg"
                        prop = "Khối lượng (kg)"
                        got_weight = True
                    elif pset_name == "Dimensions" and ("length" in pl or "diameter" in pl or "radius" in pl):
                        unit = "mm"
                    elif any(k in pl for k in _REBAR_WEIGHT_KEYS):
                        unit = "kg"
                        got_weight = True
                    else:
                        unit = _guess_unit(prop, length_unit)
                    records.append({
                        "ifc_type": ifc_type, "name": name, "storey": storey, "guid": guid,
                        "pset": pset_name, "prop": prop,
                        "value": round(num_val, 4), "unit": unit, "note": "",
                    })

            # Pset không cho ra trọng lượng → lấy từ thuộc tính STEP (trường hợp Revit).
            # Chỉ chạy khi CHƯA có bản ghi kg nào cho thanh này → không đếm trùng.
            if not got_weight:
                kg = _rebar_weight_from_step(el)
                if kg is not None:
                    records.append({
                        "ifc_type": ifc_type, "name": name, "storey": storey, "guid": guid,
                        "pset": "(thuộc tính STEP)", "prop": "Khối lượng (kg)",
                        "value": round(kg, 3), "unit": "kg",
                        "note": "Tính từ CrossSectionArea × BarLength × 7850",
                    })
                elif not qtos:
                    records.append({
                        "ifc_type": ifc_type, "name": name, "storey": storey, "guid": guid,
                        "pset": "—", "prop": "—", "value": None, "unit": "—",
                        "note": "Không có Qto/Pset cốt thép",
                    })
    return records


def _extract_qtos(ifc_file, discipline: str) -> list:
    """Mỗi record = 1 giá trị số của 1 element từ BaseQuantities (Qto_*) Psets."""
    disc = discipline.upper()
    if disc == "ALL":
        target_types = list({t for types in DISC_TYPES.values() for t in types})
    else:
        target_types = DISC_TYPES.get(disc, [])

    length_unit = _length_unit_symbol(ifc_file)
    records = []
    for ifc_type in target_types:
        try:
            elements = ifc_file.by_type(ifc_type, include_subtypes=False)
        except Exception:
            continue
        for el in elements:
            name   = (el.Name or "").strip() or "—"
            storey = _storey(el)
            group  = TYPE_GROUP_VN.get(ifc_type, ifc_type)
            guid   = el.GlobalId or ""

            # Lấy tất cả Psets rồi lọc BaseQuantities (Qto_*)
            try:
                all_psets = ifc_util.get_psets(el)
            except Exception:
                all_psets = {}
            qtos = {k: v for k, v in all_psets.items()
                    if (k.startswith("Qto_") or k == "BaseQuantities") and isinstance(v, dict)}

            if not qtos:
                records.append({
                    "ifc_type": ifc_type, "group": group, "name": name,
                    "storey": storey, "guid": guid,
                    "pset": "—", "prop": "—", "value": None, "unit": "—",
                    "note": "Không có BaseQuantities",
                })
                continue

            for pset_name, props in qtos.items():
                for prop, val in props.items():
                    if prop.lower() == "id" or val is None:
                        continue
                    try:
                        num_val = float(val)
                    except (TypeError, ValueError):
                        continue
                    records.append({
                        "ifc_type": ifc_type, "group": group, "name": name,
                        "storey": storey, "guid": guid,
                        "pset": pset_name, "prop": prop,
                        "value": round(num_val, 4),
                        "unit": _guess_unit(prop, length_unit),
                        "note": "",
                    })
    return records


# ── Bổ sung BaseQuantities từ hình học (--bo-sung-qto) ────────────────────────
# IFC xuất không bật "Export base quantities" (Duplex, Schependomlaan) hoặc sinh bằng
# code (mô hình phần ngầm 12/09) → _extract_qtos ra 0. Nhánh này đo hình học rồi GHI
# BaseQuantities vào mô hình, để ifc_qto / wbs_ifc_mapper / model_payment / payment_from_model
# đọc được nguyên bản. Không ghi đè cấu kiện đã có Qto từ phần mềm xuất.

_QTO_VOLUME_TYPES = {"IfcWall", "IfcWallStandardCase", "IfcSlab", "IfcBeam", "IfcColumn",
                     "IfcMember", "IfcFooting", "IfcPlate", "IfcStairFlight", "IfcRoof", "IfcRamp"}
_QTO_PLAN_AREA_TYPES = {"IfcSlab", "IfcCovering", "IfcRoof"}
_QTO_OPENING_TYPES = {"IfcDoor", "IfcWindow"}
_QTO_METHOD = "Đo từ hình học IFC (ifc_qto.py --bo-sung-qto)"


def _mesh_closed(verts, faces) -> bool:
    """Lưới kín = mọi cạnh tam giác dùng chung đúng 2 mặt. Lưới hở thì thể tích vô nghĩa."""
    import numpy as np
    e = np.sort(np.concatenate([faces[:, [0, 1]], faces[:, [1, 2]], faces[:, [2, 0]]]), axis=1)
    _, cnt = np.unique(e, axis=0, return_counts=True)
    return bool((cnt == 2).all())


def _bo_sung_qto(ifc_file, discipline: str) -> dict:
    """Ghi BaseQuantities đo từ hình học cho cấu kiện chưa có Qto. Trả về thống kê.

    - Thể tích (NetVolume): chỉ khi Body là khối đặc (không phải SurfaceModel), lưới kín,
      và 0 < V ≤ thể tích hộp bao. Duplex: 13 tấm trần SurfaceModel cho tổng V ~11 lần tổng hộp bao (tấm tệ nhất 36 lần).
    - Diện tích mặt bằng (NetArea) cho sàn/trần/mái: tổng hình chiếu |z| các tam giác ÷ 2
      (mặt trên + mặt dưới) — không phụ thuộc chiều pháp tuyến, cần lưới kín.
    - Cửa/cửa sổ (Area): OverallWidth × OverallHeight, không cần hình học.
    Giá trị ghi theo đơn vị project (đọc từ IFC), giống quy ước của phần mềm xuất.
    """
    import numpy as np
    import ifcopenshell.api
    import ifcopenshell.geom
    import ifcopenshell.util.shape as shape_util
    import ifcopenshell.util.unit as unit_util

    disc = discipline.upper()
    if disc == "ALL":
        target_types = list({t for types in DISC_TYPES.values() for t in types})
    else:
        target_types = DISC_TYPES.get(disc, [])

    ls = unit_util.calculate_unit_scale(ifc_file, "LENGTHUNIT")   # đơn vị project → m
    as_ = unit_util.calculate_unit_scale(ifc_file, "AREAUNIT")
    vs = unit_util.calculate_unit_scale(ifc_file, "VOLUMEUNIT")
    settings = ifcopenshell.geom.settings()                        # kết quả hình học theo m

    stats = {"da_co": 0, "the_tich": 0, "dien_tich": 0, "cua": 0, "bo_qua": defaultdict(int)}
    for ifc_type in target_types:
        try:
            elements = ifc_file.by_type(ifc_type, include_subtypes=False)
        except Exception:                                          # loại không có trong schema (IFC2x3)
            continue
        for el in elements:
            psets = ifc_util.get_psets(el)
            if any(k.startswith("Qto_") or k == "BaseQuantities" for k in psets):
                stats["da_co"] += 1
                continue

            props = {}
            if ifc_type in _QTO_OPENING_TYPES:
                w, h = getattr(el, "OverallWidth", None), getattr(el, "OverallHeight", None)
                if w and h:
                    props["Area"] = round(w * h * ls * ls / as_, 4)
                    stats["cua"] += 1
                else:
                    stats["bo_qua"]["cửa thiếu OverallWidth/Height"] += 1
            elif ifc_type in _QTO_VOLUME_TYPES | _QTO_PLAN_AREA_TYPES:
                if not el.Representation:
                    stats["bo_qua"]["không có hình học"] += 1
                    continue
                try:
                    sh = ifcopenshell.geom.create_shape(settings, el)  # giữ biến sh: viết gọn 1 dòng thì verts rỗng
                except Exception:
                    stats["bo_qua"]["lỗi dựng hình học"] += 1
                    continue
                geo = sh.geometry
                if not len(geo.verts):
                    stats["bo_qua"]["hình học rỗng (cấu kiện gộp)"] += 1
                    continue
                verts = np.array(geo.verts).reshape(-1, 3)
                faces = np.array(geo.faces).reshape(-1, 3)
                closed = _mesh_closed(verts, faces)
                surface = any(r.RepresentationType == "SurfaceModel"
                              for r in el.Representation.Representations)

                if ifc_type in _QTO_VOLUME_TYPES:
                    v = shape_util.get_volume(geo)
                    bbox = shape_util.get_x(geo) * shape_util.get_y(geo) * shape_util.get_z(geo)
                    if closed and not surface and 0 < v <= bbox * 1.001:
                        props["NetVolume"] = round(v / vs, 4)
                        stats["the_tich"] += 1
                    else:
                        reason = ("bề mặt (SurfaceModel), không phải khối" if surface else
                                  "lưới hở" if not closed else "thể tích vượt hộp bao")
                        stats["bo_qua"][f"thể tích: {reason}"] += 1

                if ifc_type in _QTO_PLAN_AREA_TYPES:
                    if closed:
                        a, b, c = verts[faces[:, 0]], verts[faces[:, 1]], verts[faces[:, 2]]
                        plan = float(np.abs(np.cross(b - a, c - a)[:, 2]).sum() / 4)
                        props["NetArea"] = round(plan / as_, 4)
                        stats["dien_tich"] += 1
                    else:
                        stats["bo_qua"]["diện tích: lưới hở"] += 1
            else:
                stats["bo_qua"][f"chưa hỗ trợ {ifc_type}"] += 1

            if props:
                qto = ifcopenshell.api.run("pset.add_qto", ifc_file, product=el, name="BaseQuantities")
                ifcopenshell.api.run("pset.edit_qto", ifc_file, qto=qto, properties=props)
                qto.MethodOfMeasurement = _QTO_METHOD
    return stats


# ── Excel export ──────────────────────────────────────────────────────────────

def _sheet_chi_tiet(wb, records, ifc_path, discipline, stage, project):
    ws = wb.active
    ws.title = "Chi_Tiet"
    headers = ["STT", "Nhóm", "IFC Type", "Tên Element",
               "Tầng", "Pset QTO", "Đại lượng", "Giá trị", "ĐVT", "Ghi chú"]
    _col_widths(ws, [5, 16, 24, 36, 14, 32, 28, 12, 8, 30])

    total_el = len({r["guid"] for r in records if r["guid"]})
    total_q  = sum(1 for r in records if r["value"] is not None)

    ws.merge_cells("A1:J1")
    c = ws.cell(1, 1, f"BÓC KHỐI LƯỢNG IFC — {project.upper()} | {discipline} | {stage}  |  {date.today().strftime('%d/%m/%Y')}")
    c.fill = _H_FILL; c.font = Font(bold=True, color="FFFFFF", size=13)
    c.alignment = _center; c.border = _border; ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:J2")
    c2 = ws.cell(2, 1, f"File: {ifc_path.name}  |  {total_el} elements  |  {total_q} giá trị số")
    c2.font = Font(italic=True, size=9); c2.alignment = _center

    for j, h in enumerate(headers, 1):
        _c(ws, 3, j, h, fill=_S_FILL, font=_H_FONT)
    ws.row_dimensions[3].height = 22
    ws.freeze_panes = "A4"

    for i, r in enumerate(records, 1):
        alt = _ALT if i % 2 == 0 else PatternFill()
        no_qto_fill = _W_FILL if r["note"] else PatternFill()
        _c(ws, i+3, 1,  i,              font=_S_FONT)
        _c(ws, i+3, 2,  r["group"],     font=_N_FONT, align=_left, fill=alt)
        _c(ws, i+3, 3,  r["ifc_type"],  font=_S_FONT, align=_left)
        _c(ws, i+3, 4,  r["name"],      font=_N_FONT, align=_left, fill=alt)
        _c(ws, i+3, 5,  r["storey"],    font=_N_FONT)
        _c(ws, i+3, 6,  r["pset"],      font=_S_FONT, align=_left)
        _c(ws, i+3, 7,  r["prop"],      font=_S_FONT, align=_left)
        _c(ws, i+3, 8,  r["value"],     font=_B_FONT if r["value"] else _N_FONT)
        _c(ws, i+3, 9,  r["unit"],      font=_N_FONT)
        _c(ws, i+3, 10, r["note"],      font=_S_FONT, align=_left, fill=no_qto_fill)


def _sheet_tong_hop(wb, records, discipline, project):
    ws = wb.create_sheet("Tong_Hop")
    headers = ["STT", "Nhóm", "Đại lượng", "ĐVT",
               "Số element", "Tổng", "Trung bình", "Min", "Max"]
    _col_widths(ws, [5, 18, 30, 8, 12, 14, 14, 12, 12])

    ws.merge_cells("A1:I1")
    t = ws.cell(1, 1, f"TỔNG HỢP KHỐI LƯỢNG — {project.upper()} | {discipline}")
    t.fill = _H_FILL; t.font = Font(bold=True, color="FFFFFF", size=13)
    t.alignment = _center; t.border = _border; ws.row_dimensions[1].height = 30

    for j, h in enumerate(headers, 1):
        _c(ws, 2, j, h, fill=_S_FILL, font=_H_FONT)
    ws.freeze_panes = "A3"

    # Nhóm theo (nhóm, prop, unit)
    summary = defaultdict(lambda: {"values": [], "guids": set()})
    for r in records:
        if r["value"] is not None:
            key = (r["group"], r["prop"], r["unit"])
            summary[key]["values"].append(r["value"])
            summary[key]["guids"].add(r["guid"])

    for stt, ((grp, prop, unit), data) in enumerate(sorted(summary.items()), 1):
        vals  = data["values"]
        n_el  = len(data["guids"])
        total = round(sum(vals), 3)
        avg   = round(sum(vals) / len(vals), 3) if vals else 0
        alt   = _ALT if stt % 2 == 0 else PatternFill()
        _c(ws, stt+2, 1, stt,   font=_N_FONT)
        _c(ws, stt+2, 2, grp,   font=_B_FONT, align=_left, fill=alt)
        _c(ws, stt+2, 3, prop,  font=_N_FONT, align=_left)
        _c(ws, stt+2, 4, unit,  font=_N_FONT)
        _c(ws, stt+2, 5, n_el,  font=_N_FONT)
        _c(ws, stt+2, 6, total, font=_B_FONT)
        _c(ws, stt+2, 7, avg,   font=_N_FONT)
        _c(ws, stt+2, 8, round(min(vals), 3), font=_N_FONT)
        _c(ws, stt+2, 9, round(max(vals), 3), font=_N_FONT)

    return summary


def _sheet_tom_tat_boq(wb, summary, discipline, stage, project):
    ws = wb.create_sheet("Tom_Tat_BoQ")
    _col_widths(ws, [5, 22, 30, 8, 14, 14])

    ws.merge_cells("A1:F1")
    t = ws.cell(1, 1, f"TÓM TẮT KL THEO NHÓM — {project.upper()} | {discipline} | {stage}")
    t.fill = _H_FILL; t.font = Font(bold=True, color="FFFFFF", size=13)
    t.alignment = _center; t.border = _border; ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:F2")
    n = ws.cell(2, 1,
        "ℹ Đơn vị dài (Length/Width/Height) đọc tự động từ project unit của IFC "
        "(Revit thường mm); Area/Volume theo m²/m³. Cốt thép quy đổi kg (7850 kg/m³).")
    n.font = Font(italic=True, color="C00000", size=9); n.alignment = _left

    for j, h in enumerate(["STT", "Nhóm", "Đại lượng chính", "ĐVT", "Tổng KL", "Số element"], 1):
        _c(ws, 3, j, h, fill=_S_FILL, font=_H_FONT)
    ws.freeze_panes = "A4"

    # Chọn 1 đại lượng đại diện mỗi nhóm: Volume > Area > Length
    best: dict = {}
    for (grp, prop, unit), data in summary.items():
        score = _priority(prop)
        vals  = data["values"]
        if grp not in best or score < best[grp]["score"]:
            best[grp] = {
                "prop": prop, "unit": unit,
                "total": round(sum(vals), 3),
                "n": len(data["guids"]),
                "score": score,
            }

    for stt, (grp, info) in enumerate(sorted(best.items()), 1):
        alt = _ALT if stt % 2 == 0 else PatternFill()
        _c(ws, stt+3, 1, stt,            font=_N_FONT)
        _c(ws, stt+3, 2, grp,            font=_B_FONT, align=_left, fill=alt)
        _c(ws, stt+3, 3, info["prop"],   font=_N_FONT, align=_left)
        _c(ws, stt+3, 4, info["unit"],   font=_N_FONT)
        _c(ws, stt+3, 5, info["total"],  font=Font(bold=True, size=11))
        _c(ws, stt+3, 6, info["n"],      font=_N_FONT)


def _sheet_cot_thep(wb, rebar_records, project):
    """Sheet cốt thép: chỉ tạo khi có entity IfcReinforcingBar/Mesh/Tendon."""
    ws = wb.create_sheet("Cot_Thep")
    headers = ["STT", "Loại cốt thép", "Tên", "Tầng", "Pset", "Đại lượng", "Giá trị", "ĐVT", "Ghi chú"]
    _col_widths(ws, [5, 22, 36, 14, 30, 28, 12, 8, 36])

    ws.merge_cells("A1:I1")
    _W_FILL2 = PatternFill("solid", fgColor="C00000")
    c = ws.cell(1, 1, f"CỐT THÉP — {project.upper()} (từ IFC BaseQuantities)")
    c.fill = _W_FILL2; c.font = Font(bold=True, color="FFFFFF", size=13)
    c.alignment = _center; c.border = _border; ws.row_dimensions[1].height = 30

    has_data = any(r["value"] is not None for r in rebar_records)
    ws.merge_cells("A2:I2")
    note_text = (
        "⚠ Không tìm thấy dữ liệu cốt thép số (Weight/CrossSectionArea). "
        "Cách bật trong Revit: File → Export → IFC → Modify Setup → "
        "check 'Export base quantities' + 'Include structural rebar'."
        if not has_data else
        "Đơn vị Weight: kg. Kiểm tra lại IFC Units nếu giá trị bất thường."
    )
    n = ws.cell(2, 1, note_text)
    n.font = Font(italic=True, size=9, color="C00000" if not has_data else "808080")
    n.alignment = _left

    for j, h in enumerate(headers, 1):
        _c(ws, 3, j, h, fill=_S_FILL, font=_H_FONT)
    ws.row_dimensions[3].height = 22
    ws.freeze_panes = "A4"

    for i, r in enumerate(rebar_records, 1):
        alt = _ALT if i % 2 == 0 else PatternFill()
        no_data = _W_FILL if r["note"] else PatternFill()
        _c(ws, i+3, 1, i,             font=_S_FONT)
        _c(ws, i+3, 2, r["ifc_type"], font=_S_FONT, align=_left)
        _c(ws, i+3, 3, r["name"],     font=_N_FONT, align=_left, fill=alt)
        _c(ws, i+3, 4, r["storey"],   font=_N_FONT)
        _c(ws, i+3, 5, r["pset"],     font=_S_FONT, align=_left)
        _c(ws, i+3, 6, r["prop"],     font=_S_FONT, align=_left)
        _c(ws, i+3, 7, r["value"],    font=_B_FONT if r["value"] else _N_FONT)
        _c(ws, i+3, 8, r["unit"],     font=_N_FONT)
        _c(ws, i+3, 9, r["note"],     font=_S_FONT, align=_left, fill=no_data)


def export_excel(records, rebar_records, ifc_path, discipline, stage, project, out_path):
    wb = openpyxl.Workbook()
    _sheet_chi_tiet(wb, records, ifc_path, discipline, stage, project)
    summary = _sheet_tong_hop(wb, records, discipline, project)
    _sheet_tom_tat_boq(wb, summary, discipline, stage, project)
    if rebar_records is not None:
        _sheet_cot_thep(wb, rebar_records, project)
    wb.save(str(out_path))
    print(f"  ✅ Excel: {out_path.name}")


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Bóc khối lượng (QTO) từ file IFC")
    parser.add_argument("--ifc",        required=True,
                        help="File IFC (path tuyệt đối hoặc tương đối từ thư mục hiện tại)")
    parser.add_argument("--discipline", default="ALL",
                        choices=["KT", "KC", "MEP", "ALL"],
                        help="Discipline: KT / KC / MEP / ALL (mặc định: ALL)")
    parser.add_argument("--stage",      default="TKKT",
                        choices=["TKCS", "TKKT", "TKTC", "AS-BUILT"],
                        help="Giai đoạn (dùng để ghi vào tiêu đề)")
    parser.add_argument("--project",    default="",
                        help="Tên dự án (mặc định: lấy từ tên file IFC)")
    parser.add_argument("--bo-sung-qto", action="store_true",
                        help="IFC thiếu BaseQuantities: đo từ hình học, ghi bản sao "
                             "<tên>_BoSungQTO.ifc vào thư mục kết quả rồi bóc KL trên bản đó")
    args = parser.parse_args()

    ifc_path = Path(args.ifc)
    if not ifc_path.is_absolute():
        ifc_path = BASE / args.ifc
    if not ifc_path.exists():
        print(f"Lỗi: không tìm thấy file IFC: {ifc_path}")
        sys.exit(1)

    project = args.project or ifc_path.stem
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    size_mb = ifc_path.stat().st_size // 1_048_576

    print(f"\n{'='*60}")
    print(f"  IFC QTO — {project}")
    print(f"  File       : {ifc_path.name}  ({size_mb} MB)")
    print(f"  Discipline : {args.discipline}  |  Stage: {args.stage}")
    print(f"{'='*60}")
    print("  Đang mở và đọc IFC...")

    ifc_file = ifcopenshell.open(str(ifc_path))

    if args.bo_sung_qto:
        print("  Đang đo hình học để bổ sung BaseQuantities...")
        st = _bo_sung_qto(ifc_file, args.discipline)
        qto_ifc = OUTPUT_DIR / f"{ifc_path.stem}_BoSungQTO.ifc"
        ifc_file.write(str(qto_ifc))
        print(f"    Đã có Qto sẵn (giữ nguyên) : {st['da_co']}")
        print(f"    Ghi thể tích (NetVolume)   : {st['the_tich']}")
        print(f"    Ghi diện tích (NetArea)    : {st['dien_tich']}")
        print(f"    Ghi diện tích cửa (Area)   : {st['cua']}")
        for reason, n in sorted(st["bo_qua"].items()):
            print(f"    Bỏ qua — {reason}: {n}")
        print(f"  ✅ IFC đã bổ sung: {qto_ifc}")
        ifc_path = qto_ifc

    print("  Đang bóc khối lượng từ BaseQuantities Psets...")
    records = _extract_qtos(ifc_file, args.discipline)

    if not records:
        print("  Không tìm thấy element nào. Thử --discipline ALL\n")
        sys.exit(0)

    total_el = len({r["guid"] for r in records if r["guid"]})
    total_q  = sum(1 for r in records if r["value"] is not None)
    no_qto   = sum(1 for r in records if r["note"] == "Không có BaseQuantities")

    # Cốt thép — chỉ chạy với KC hoặc ALL
    disc_upper = args.discipline.upper()
    rebar_records = None
    if disc_upper in ("KC", "ALL"):
        print("  Đang tìm dữ liệu cốt thép (IfcReinforcingBar/Mesh/Tendon)...")
        rebar_records = _extract_rebar(ifc_file, args.discipline)
        n_rebar_el = len({r["guid"] for r in rebar_records if r["guid"]})
        n_rebar_val = sum(1 for r in rebar_records if r["value"] is not None)
        if n_rebar_el == 0:
            print("  ⚠  Không tìm thấy entity cốt thép trong IFC.")
            print("     Nếu mô hình có cốt thép, cần export Revit với:")
            print("       File → Export → IFC → Modify Setup → check 'Include structural rebar'")
        else:
            print(f"  → Cốt thép: {n_rebar_el} entities, {n_rebar_val} giá trị số")

    print(f"\n  Kết quả:")
    print(f"    Tổng elements     : {total_el}")
    print(f"    Có BaseQuantities : {total_el - no_qto}")
    print(f"    Không có Qto      : {no_qto}")
    print(f"    Số giá trị số     : {total_q}")

    if no_qto == total_el:
        print("\n  ⚠  Toàn bộ elements không có BaseQuantities.")
        print("     Cách bật trong Revit:")
        print("       File → Export → IFC → Modify Setup → check 'Export base quantities'")
        print("     Script vẫn xuất Excel, cột Ghi chú sẽ hiển thị cảnh báo.")

    slug     = re.sub(r"[^\w]", "_", project)
    out_path = OUTPUT_DIR / f"IFC_QTO_{slug}_{args.discipline}_{date.today().strftime('%Y%m%d')}.xlsx"

    print("\n  Đang xuất báo cáo...")
    export_excel(records, rebar_records, ifc_path, args.discipline, args.stage, project, out_path)

    sheets = "Chi_Tiet | Tong_Hop | Tom_Tat_BoQ"
    if rebar_records is not None:
        sheets += " | Cot_Thep"
    print(f"\n  ✅ Hoàn thành! Kết quả tại: {out_path}")
    print(f"     Sheets: {sheets}")
    if no_qto == total_el:
        print("\n  ⚠  Báo cáo xuất nhưng chưa có số liệu KL — xem hướng dẫn bên trên.")
    print()


if __name__ == "__main__":
    main()
