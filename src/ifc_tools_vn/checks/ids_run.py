"""
ids_run.py — Chấm mô hình IFC theo file IDS chuẩn buildingSMART (engine ifctester).
================================================================================
Khác ids_checker.py cũ: KHÔNG dùng luật tự chế. Thước đo là file .ids (đặc tả hợp đồng).
Engine chấm là ifctester chính thức của IfcOpenShell → kết quả có thể tái lập ở bất kỳ
công cụ IDS nào, khách không cãi được.

Xuất 3 loại báo cáo:
  1. Excel tiếng Việt  — tóm tắt compliance + chi tiết lỗi (cho hồ sơ VN)
  2. HTML chuẩn        — bản của chính ifctester (đối chứng độc lập)
  3. BCF               — mở thẳng issue trên viewer 3D;
                         đã chuẩn hoá cho BIMcollab Zoom (_bcf_cho_zoom)

Cách dùng:
    python ids_run.py --ids EIR_..._KC_TKTC.ids --ifc Duplex_A_20110907.ifc --project Duplex
"""
import argparse
import sys
from pathlib import Path
from datetime import date

sys.stdout.reconfigure(encoding="utf-8")

try:
    import ifcopenshell
    from ifctester import ids, reporter
except ImportError as e:
    print(f"Lỗi import: {e}\n-> python -m pip install ifctester ifcopenshell")
    sys.exit(1)

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Lỗi: pip install openpyxl")
    sys.exit(1)

# Đo độ phủ (module dùng chung, cùng thư mục) — vạch ra 3 kiểu "100% đạt mà chưa kiểm gì".
sys.path.insert(0, str(Path(__file__).parent))
from ids_dophu import phan_tich, dong_canh_bao, co_canh_bao

from ifc_tools_vn.config import BASE, OUTPUT_DIR as OUTDIR  # đường dẫn: xem config.py

# ── styles ──
_H = PatternFill("solid", fgColor="1F4E79")
_S = PatternFill("solid", fgColor="2E75B6")
_OK = PatternFill("solid", fgColor="E2EFDA")
_ER = PatternFill("solid", fgColor="FCE4D6")
_KA = PatternFill("solid", fgColor="FFD966")   # vàng hổ phách: KHÔNG kiểm được (0 cấu kiện áp dụng)
_ALT = PatternFill("solid", fgColor="DEEAF1")
_thin = Side(style="thin", color="CCCCCC")
_bd = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
_lft = Alignment(horizontal="left", vertical="center", wrap_text=True)
_HF = Font(bold=True, color="FFFFFF", size=10)
_BF = Font(bold=True, size=10)
_NF = Font(size=10)
_SF = Font(size=9)


def _c(ws, r, col, v=None, fill=None, font=None, align=None):
    cell = ws.cell(row=r, column=col, value=v)
    cell.border, cell.fill = _bd, fill or PatternFill()
    cell.font, cell.alignment = font or _NF, align or _ctr
    return cell


def _widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _storey(element) -> str:
    """Tên IfcBuildingStorey chứa element — port từ ids_checker.py (bộ cũ)."""
    for rel in getattr(element, "ContainedInStructure", None) or []:
        container = rel.RelatingStructure
        if container.is_a("IfcBuildingStorey"):
            return container.Name or ""
    return ""


def export_excel(data, ifc_name, project, out_path, model=None, dophu=None):
    wb = openpyxl.Workbook()

    # Sheet 1 — Tổng hợp theo spec
    ws = wb.active
    ws.title = "TongHop_Spec"
    _widths(ws, [7, 16, 46, 10, 10, 10, 10])
    ws.merge_cells("A1:G1")
    t = ws.cell(1, 1, f"KIỂM TRA IDS (chuẩn buildingSMART) — {project.upper()}  |  {date.today().strftime('%d/%m/%Y')}")
    t.fill, t.font, t.alignment, t.border = _H, Font(bold=True, color="FFFFFF", size=13), _ctr, _bd
    ws.row_dimensions[1].height = 28
    ws.merge_cells("A2:G2")
    pc = data["percent_specifications_pass"]
    # Dòng 2 phải nói cả "đã kiểm được gì" — nếu chỉ ghi % đạt thì một mô hình thiếu
    # hẳn cấu kiện vẫn hiện 100% (spec 0 cấu kiện áp dụng vẫn được chuẩn IDS tính là đạt).
    _canh_bao = ""
    if dophu:
        _n_ka = len(dophu["spec_khong_ap_dung"])
        _n_kp = sum(dophu["lop_khong_phu"].values())
        if _n_ka:
            _canh_bao += f"  |  ⚠ {_n_ka} yêu cầu KHÔNG kiểm được (0 cấu kiện áp dụng)"
        if _n_kp:
            _canh_bao += f"  |  ⚠ {_n_kp} cấu kiện không yêu cầu nào phủ"
        _canh_bao += f"  |  Độ phủ: {dophu['cau_kien_phu']}/{dophu['cau_kien_tong']}"
    s2 = ws.cell(2, 1, f"File: {ifc_name}  |  Compliance spec: {pc}%  "
                       f"({data['total_specifications_pass']}/{data['total_specifications']} spec đạt)  |  "
                       f"Checks: {data['total_checks_pass']}/{data['total_checks']}{_canh_bao}")
    s2.font, s2.alignment = Font(italic=True, size=9, bold=bool(_canh_bao)), _ctr
    for j, h in enumerate(["STT", "Mã EIR", "Yêu cầu", "Áp dụng", "Đạt", "Trượt", "% Đạt"], 1):
        _c(ws, 3, j, h, fill=_S, font=_HF)
    ws.freeze_panes = "A4"
    for i, sp in enumerate(data["specifications"], 1):
        code = sp.get("name", "").split(" — ")[0]
        req = sp.get("name", "")
        appl = sp.get("total_applicable", 0)
        passed = sp.get("total_applicable_pass", 0)
        failed = sp.get("total_applicable_fail", 0)
        pct = sp.get("percent_applicable_pass", 100)
        # 0 cấu kiện áp dụng => ifctester tính ĐẠT, nhưng KHÔNG được tô xanh:
        # không kiểm gì thì không phải là đạt. Tô vàng hổ phách và ghi rõ.
        khong_kiem = (appl == 0)
        row_fill = _KA if khong_kiem else (_OK if sp.get("status") else _ER)
        alt = _ALT if i % 2 == 0 else PatternFill()
        _c(ws, i + 3, 1, i, font=_SF)
        _c(ws, i + 3, 2, code, font=Font(bold=True, size=9))
        _c(ws, i + 3, 3, req, font=_NF, align=_lft, fill=alt)
        _c(ws, i + 3, 4, appl, font=_NF)
        _c(ws, i + 3, 5, passed, font=_NF, fill=_OK if (failed == 0 and not khong_kiem) else PatternFill())
        _c(ws, i + 3, 6, failed, font=_NF, fill=_ER if failed else PatternFill())
        _c(ws, i + 3, 7, "KHÔNG KIỂM" if khong_kiem else f"{pct}%", font=_BF, fill=row_fill)

    # Sheet 2 — Chi tiết lỗi
    ws2 = wb.create_sheet("Loi_Chi_Tiet")
    _widths(ws2, [6, 16, 20, 34, 24, 14, 44])
    ws2.merge_cells("A1:G1")
    t2 = ws2.cell(1, 1, f"CHI TIẾT LỖI — {project.upper()}")
    t2.fill, t2.font, t2.alignment, t2.border = PatternFill("solid", fgColor="C00000"), Font(bold=True, color="FFFFFF", size=13), _ctr, _bd
    ws2.row_dimensions[1].height = 26
    for j, h in enumerate(["STT", "Mã EIR", "IFC Type", "Tên cấu kiện", "GUID", "Tầng", "Lý do trượt"], 1):
        _c(ws2, 2, j, h, fill=_S, font=_HF)
    ws2.freeze_panes = "A3"
    row = 3
    n = 0
    for sp in data["specifications"]:
        if sp.get("status"):
            continue
        code = sp.get("name", "").split(" — ")[0]
        for req in sp.get("requirements", []):
            for fe in req.get("failed_entities", []):
                n += 1
                guid = fe.get("global_id") or ""
                storey = "—"
                if model is not None and guid:
                    try:
                        storey = _storey(model.by_guid(guid)) or "—"
                    except Exception:
                        storey = "—"
                _c(ws2, row, 1, n, font=_SF)
                _c(ws2, row, 2, code, font=Font(bold=True, size=9))
                _c(ws2, row, 3, fe.get("class") or "—", font=_SF, align=_lft)
                _c(ws2, row, 4, (fe.get("name") or "—"), font=_NF, align=_lft, fill=_ER)
                _c(ws2, row, 5, guid or "—", font=_SF, align=_lft)
                _c(ws2, row, 6, storey, font=_NF)
                _c(ws2, row, 7, fe.get("reason") or "—", font=_SF, align=_lft, fill=_ER)
                row += 1
    if n == 0:
        ws2.merge_cells(f"A3:G3")
        # Không có lỗi KHÔNG đồng nghĩa đã kiểm hết — xem sheet DoPhu_CanhBao.
        if dophu and co_canh_bao(dophu):
            _c(ws2, 3, 1, "⚠ Không phát hiện lỗi trên phần ĐÃ kiểm — nhưng chưa kiểm hết. "
                          "Xem sheet 'DoPhu_CanhBao'.", fill=_KA, font=_BF, align=_lft)
        else:
            _c(ws2, 3, 1, "✅ Không có lỗi — mô hình đạt toàn bộ yêu cầu IDS.", fill=_OK, font=_BF, align=_lft)

    # ── Sheet 3 — Độ phủ & cảnh báo (thêm mới, không đụng 2 sheet trên) ──
    if dophu:
        ws3 = wb.create_sheet("DoPhu_CanhBao")
        _widths(ws3, [6, 34, 16, 62])
        ws3.merge_cells("A1:D1")
        t3 = ws3.cell(1, 1, f"ĐÃ KIỂM ĐƯỢC NHỮNG GÌ — {project.upper()}")
        t3.fill, t3.font, t3.alignment, t3.border = _H, Font(bold=True, color="FFFFFF", size=13), _ctr, _bd
        ws3.row_dimensions[1].height = 26
        r = 2
        for line in dong_canh_bao(dophu):
            ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            _c(ws3, r, 1, line.strip(), font=_NF, align=_lft,
               fill=_KA if line.strip().startswith(("⚠", "⛔")) else PatternFill())
            r += 1
        r += 1
        for j, h in enumerate(["STT", "Lớp IFC không yêu cầu nào phủ", "Số cấu kiện", "Ghi chú"], 1):
            _c(ws3, r, j, h, fill=_S, font=_HF)
        r += 1
        if dophu["lop_khong_phu"]:
            for i, (cls, cnt) in enumerate(dophu["lop_khong_phu"].items(), 1):
                _c(ws3, r, 1, i, font=_SF)
                _c(ws3, r, 2, cls, font=_NF, align=_lft, fill=_KA)
                _c(ws3, r, 3, cnt, font=_BF)
                _c(ws3, r, 4, "Không spec nào áp dụng lên lớp này — cân nhắc bổ sung yêu cầu vào EIR",
                   font=_SF, align=_lft)
                r += 1
        else:
            ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            _c(ws3, r, 1, "— Mọi cấu kiện trong mô hình đều có yêu cầu phủ.", fill=_OK, font=_NF, align=_lft)
            r += 1
        r += 1
        for j, h in enumerate(["STT", "Mã EIR", "Tình trạng", "Yêu cầu không kiểm được"], 1):
            _c(ws3, r, j, h, fill=_S, font=_HF)
        r += 1
        chet = {ma for ma, _, _ in dophu["spec_luat_chet"]}
        if dophu["spec_khong_ap_dung"]:
            for i, (ma, ten) in enumerate(dophu["spec_khong_ap_dung"], 1):
                _c(ws3, r, 1, i, font=_SF)
                _c(ws3, r, 2, ma, font=Font(bold=True, size=9))
                _c(ws3, r, 3, "LUẬT CHẾT" if ma in chet else "0 cấu kiện", font=_BF, fill=_KA)
                _c(ws3, r, 4, ten, font=_NF, align=_lft)
                r += 1
        else:
            ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            _c(ws3, r, 1, "— Mọi yêu cầu đều tìm được cấu kiện để kiểm.", fill=_OK, font=_NF, align=_lft)
    wb.save(str(out_path))


def _png_nen(w=160, h=100):
    """Ảnh nền cho snapshot — Zoom cần có ảnh, ifctester không tự chụp được hình."""
    import struct
    import zlib
    dong = []
    for y in range(h):
        d = bytearray([0])
        for x in range(w):
            vien = x < 3 or y < 3 or x >= w - 3 or y >= h - 3
            d += bytes((220, 38, 38) if vien else (238, 242, 247))
        dong.append(bytes(d))
    def khoi(t, b):
        return struct.pack(">I", len(b)) + t + b + struct.pack(">I", zlib.crc32(t + b) & 0xFFFFFFFF)
    return (b"\x89PNG\r\n\x1a\n" + khoi(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
            + khoi(b"IDAT", zlib.compress(b"".join(dong))) + khoi(b"IEND", b""))


def _bcf_cho_zoom(duong_dan, model, data):
    """Chuẩn hoá BCF của ifctester để BIMcollab Zoom mở được (đã thử thật trên Zoom 9.9, 04/10/2026).

    BCF gốc đúng chuẩn 2.1 nhưng Zoom báo "no viewpoint": Zoom cần đúng khuôn file của nó —
    `viewpoint.bcfv` + `snapshot.png` + 1 bình luận trỏ góc nhìn. Camera gốc nhìn chéo từ xa 17 m
    nên tường bên trong bị mái che → đặt lại camera theo khung bao thật của cấu kiện, nhìn dốc
    xuống, kèm mặt cắt ngang trong cấu kiện, dưới đáy sàn/mái/trần trùm lên nó
    (Zoom hiểu Direction +Z = bỏ phần phía trên).
    Phía Zoom phải chọn Zoom-to "Full" thì mới tô sáng cấu kiện (chế độ Quick bỏ qua Selection).
    Số thứ tự lỗi (Index) = cột STT sheet Loi_Chi_Tiet → lỗi N trong Zoom là dòng N trong Excel
    (ifctester sinh mã topic ngẫu nhiên nên nếu không gán thì số bị xáo mỗi lần chạy).
    Trả về (số lỗi, số camera đặt theo cấu kiện).
    """
    import re
    import uuid
    import zipfile
    from xml.sax.saxutils import unescape
    import ifcopenshell.geom

    st = ifcopenshell.geom.settings()
    st.set("use-world-coords", True)

    def khung(guid):
        try:
            sh = ifcopenshell.geom.create_shape(st, model.by_guid(guid))  # giữ biến sh: viết gọn 1 dòng thì verts rỗng
            v = sh.geometry.verts
        except Exception:
            return None   # cấu kiện không có hình (vd IfcBuilding) → giữ camera gốc
        if not v:
            return None
        xs, ys, zs = v[0::3], v[1::3], v[2::3]
        return (min(xs), min(ys), min(zs)), (max(xs), max(ys), max(zs))

    _che = {}
    def tam_che():
        """Khung bao sàn/mái/trần — chỉ tính 1 lần, khi có lỗi cần đặt mặt cắt."""
        if "ds" not in _che:
            _che["ds"] = [k for e in model.by_type("IfcSlab") + model.by_type("IfcRoof") + model.by_type("IfcCovering")
                          if (k := khung(e.GlobalId))]
        return _che["ds"]

    # Thứ tự dòng STT của sheet Loi_Chi_Tiet (cùng vòng lặp với export_excel).
    stt = [(fe.get("global_id") or "", sp.get("name", ""))
           for sp in data["specifications"] if not sp.get("status")
           for req in sp.get("requirements", []) for fe in req.get("failed_entities", [])]
    da_gan = set()

    def so_stt(guid, mo_ta):
        for i, (g, ten_spec) in enumerate(stt):
            if i not in da_gan and g == guid and mo_ta.startswith(ten_spec):
                da_gan.add(i)
                return i + 1
        return None

    xyz = lambda v: "".join(f"<{k}>{x:.4f}</{k}>" for k, x in zip("XYZ", v))
    anh = _png_nen()
    goc = zipfile.ZipFile(duong_dan)
    tam = duong_dan.with_suffix(".tmp")
    so_loi = so_cam = 0
    with zipfile.ZipFile(tam, "w", zipfile.ZIP_DEFLATED) as ra:
        ra.writestr("bcf.version", '<?xml version="1.0" encoding="UTF-8"?>\n<Version VersionId="2.1">\n'
                                   "  <DetailedVersion>2.1</DetailedVersion>\n</Version>\n")
        ra.writestr("project.bcfp", goc.read("project.bcfp"))
        for ten in sorted(n for n in goc.namelist() if n.endswith("markup.bcf")):
            thu_muc = ten.split("/")[0]
            mk = goc.read(ten).decode("utf-8")
            vps = re.findall(r'<Viewpoints Guid="([^"]+)">', mk)
            if not vps:
                continue
            g_vp = vps[0]   # ifctester ghi 2 góc nhìn trùng nhau → giữ 1
            so_loi += 1
            vp = goc.read(f"{thu_muc}/{g_vp}.bcfv").decode("utf-8")
            vp = vp.replace("<Components>", '<Components>\n    <ViewSetupHints SpacesVisible="false" '
                                            'SpaceBoundariesVisible="false" OpeningsVisible="false"/>', 1)
            m = re.search(r'<Selection>\s*<Component IfcGuid="([^"]+)"', vp)
            k = khung(m.group(1)) if m else None
            if k:
                lo, hi = k
                c = [(a + b) / 2 for a, b in zip(lo, hi)]
                cheo = sum((b - a) ** 2 for a, b in zip(lo, hi)) ** 0.5
                n = (0.25 + 0.25 + 1.0) ** 0.5
                d = (-0.5 / n, -0.5 / n, -1.0 / n)          # nhìn dốc xuống ~55°
                p = [ci - di * max(5.0, 2.0 * cheo) for ci, di in zip(c, d)]
                up = (-d[2] * d[0], -d[2] * d[1], 1 - d[2] * d[2])
                mu = sum(x * x for x in up) ** 0.5
                up = [x / mu for x in up]
                vp = re.sub(r"<PerspectiveCamera>.*?</PerspectiveCamera>",
                            f"<PerspectiveCamera><CameraViewPoint>{xyz(p)}</CameraViewPoint>"
                            f"<CameraDirection>{xyz(d)}</CameraDirection><CameraUpVector>{xyz(up)}</CameraUpVector>"
                            "<FieldOfView>60</FieldOfView></PerspectiveCamera>", vp, flags=re.S)
                # Mặt cắt luôn nằm TRONG cấu kiện và thấp hơn đáy mọi sàn/mái/trần trùm lên nó: cắt trên đỉnh
                # thì sàn tầng trên / sàn nền / trần vẫn che (gặp thật 04/10: biệt thự 35/43 lỗi, Duplex 48/57).
                z_cat = min(lo[2] + 1.2, hi[2] - 0.02)          # mặc định cao 1,2 m như mặt bằng Revit
                for a, b in tam_che():
                    if a[0] < hi[0] and b[0] > lo[0] and a[1] < hi[1] and b[1] > lo[1] and lo[2] + 0.1 < a[2] < z_cat:
                        z_cat = a[2] - 0.02
                vp = vp.replace("</VisualizationInfo>",
                                f"<ClippingPlanes><ClippingPlane><Location>{xyz((c[0], c[1], z_cat))}</Location>"
                                "<Direction><X>0</X><Y>0</Y><Z>1</Z></Direction></ClippingPlane></ClippingPlanes>\n"
                                "</VisualizationInfo>")
                so_cam += 1
            ma = re.search(r"<Description>([A-Z]+-[A-Z]+-\d+)", mk)
            ngay = re.search(r"<CreationDate>(.*?)</CreationDate>", mk)
            mk = re.sub(r"\s*<Viewpoints Guid=.*?</Viewpoints>", "", mk, flags=re.S)
            mk = re.sub(r'TopicType="[^"]*"', 'TopicType="Issue"', mk, count=1)
            mk = re.sub(r'TopicStatus="[^"]*"', 'TopicStatus="Active"', mk, count=1)
            mo_ta = re.search(r"<Description>(.*?)</Description>", mk, re.S)
            so = so_stt(m.group(1) if m else "", unescape(mo_ta.group(1)) if mo_ta else "") or len(stt) + so_loi
            if "<Index>" not in mk:
                mk = mk.replace("<CreationDate>", f"<Index>{so}</Index>\n    <CreationDate>", 1)
            mk = mk.replace("</Markup>",
                            f'  <Comment Guid="{uuid.uuid4()}">\n    <Date>{ngay.group(1) if ngay else ""}</Date>\n'
                            f"    <Author>IfcTester</Author>\n    <Comment>{ma.group(1) if ma else 'IDS'}</Comment>\n"
                            f'    <Viewpoint Guid="{g_vp}"/>\n  </Comment>\n'
                            f'  <Viewpoints Guid="{g_vp}">\n    <Viewpoint>viewpoint.bcfv</Viewpoint>\n'
                            "    <Snapshot>snapshot.png</Snapshot>\n  </Viewpoints>\n</Markup>")
            ra.writestr(ten, mk.encode("utf-8"))
            ra.writestr(f"{thu_muc}/viewpoint.bcfv", vp.encode("utf-8"))
            ra.writestr(f"{thu_muc}/snapshot.png", anh)
    goc.close()
    tam.replace(duong_dan)
    return so_loi, so_cam


def main():
    ap = argparse.ArgumentParser(description="Chấm IFC theo IDS chuẩn (ifctester)")
    ap.add_argument("--ids", required=True, help="File .ids (đặc tả hợp đồng)")
    ap.add_argument("--ifc", required=True, help="File IFC cần chấm")
    ap.add_argument("--project", default="", help="Tên dự án (mặc định: tên file IFC)")
    ap.add_argument("--no-bcf", action="store_true", help="Bỏ xuất BCF")
    args = ap.parse_args()

    ids_path = Path(args.ids)
    ifc_path = Path(args.ifc)
    if not ifc_path.is_absolute():
        ifc_path = BASE / args.ifc
    for p, lbl in [(ids_path, "IDS"), (ifc_path, "IFC")]:
        if not p.exists():
            print(f"Lỗi: không thấy file {lbl}: {p}")
            sys.exit(1)

    project = args.project or ifc_path.stem
    OUTDIR.mkdir(parents=True, exist_ok=True)
    slug = "".join(c if c.isalnum() else "_" for c in project)

    print(f"\n{'='*60}")
    print(f"  KIỂM TRA IDS CHUẨN — {project}")
    print(f"  IDS : {ids_path.name}")
    print(f"  IFC : {ifc_path.name}")
    print(f"{'='*60}")

    doc = ids.open(str(ids_path), validate=True)   # validate schema chính thức
    model = ifcopenshell.open(str(ifc_path))
    doc.validate(model)

    data = reporter.Json(doc).report()
    # ifctester trả cấu kiện trượt theo thứ tự KHÔNG cố định giữa các lần chạy (gặp thật trên biệt thự
    # 04/10) → sắp theo tên để STT Excel và số lỗi BCF giống nhau mỗi lần chạy.
    for sp in data["specifications"]:
        for req in sp.get("requirements", []):
            req.get("failed_entities", []).sort(key=lambda fe: (fe.get("name") or "", fe.get("global_id") or ""))
    pc_spec = data["percent_specifications_pass"]
    pc_chk = data["percent_checks_pass"]
    print(f"  Compliance spec : {pc_spec}%  ({data['total_specifications_pass']}/{data['total_specifications']} spec)")
    print(f"  Compliance check: {pc_chk}%  ({data['total_checks_pass']}/{data['total_checks']} lượt kiểm)")

    # Độ phủ: trả lời "đã kiểm được những gì", vì % đạt một mình có thể là 100% mà
    # chưa kiểm gì (spec 0 cấu kiện áp dụng vẫn được chuẩn IDS tính là đạt).
    dophu = phan_tich(doc, model)

    fails = [s for s in data["specifications"] if not s.get("status")]
    if fails:
        print(f"\n  {len(fails)} yêu cầu CHƯA đạt:")
        for s in fails:
            code = s.get("name", "").split(" — ")[0]
            print(f"    ❌ {code}: {s.get('total_applicable_fail')}/{s.get('total_applicable')} cấu kiện trượt "
                  f"— {s.get('name','').split(' — ',1)[-1]}")
    elif co_canh_bao(dophu):
        print("\n  ⚠ Không phát hiện lỗi trên phần ĐÃ kiểm — nhưng chưa kiểm hết (xem dưới).")
    else:
        print("\n  ✅ Đạt toàn bộ yêu cầu IDS.")

    print(f"\n{'-'*60}\n  ĐÃ KIỂM ĐƯỢC NHỮNG GÌ\n{'-'*60}")
    for d in dong_canh_bao(dophu):
        print(d)

    # Xuất báo cáo
    xlsx = OUTDIR / f"IDS_KetQua_{slug}_{date.today().strftime('%Y%m%d')}.xlsx"
    export_excel(data, ifc_path.name, project, xlsx, model, dophu)
    html = OUTDIR / f"IDS_KetQua_{slug}_{date.today().strftime('%Y%m%d')}.html"
    html_rep = reporter.Html(doc)
    html_rep.report()          # Html (kế thừa Json) cũng cần report() trước khi ghi file
    html_rep.to_file(str(html))
    print(f"\n  📄 Excel VN : {xlsx}")
    print(f"  📄 HTML     : {html}")
    if not args.no_bcf:
        try:
            bcf = OUTDIR / f"IDS_KetQua_{slug}_{date.today().strftime('%Y%m%d')}.bcf"
            bcf_rep = reporter.Bcf(doc)
            bcf_rep.report()          # Bcf cần report() trước khi ghi file (giống Html)
            bcf_rep.to_file(str(bcf))
            n_loi, n_cam = _bcf_cho_zoom(bcf, model, data)
            print(f"  📄 BCF      : {bcf}  ({n_loi} lỗi, {n_cam} camera đặt theo cấu kiện"
                  " — mở bằng BIMcollab Zoom, Zoom-to chọn 'Full')")
        except Exception as e:
            print(f"  (bỏ qua BCF: {e})")
    print()


if __name__ == "__main__":
    main()
